#!/usr/bin/env python3
"""Generate PowerPoint slides for 70B config exploration results."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pathlib import Path

OUTPUT_DIR = Path(__file__).parent
CHART_PATH = OUTPUT_DIR / "figures" / "pareto_latency_throughput.png"
BLIS_PARETO_PATH = OUTPUT_DIR / "figures" / "blis_pareto_frontier.png"


def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # Title
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.4), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Llama 3.1 70B Config Exploration"
    p.font.size = Pt(32)
    p.font.bold = True
    p.alignment = PP_ALIGN.LEFT

    p2 = tf.add_paragraph()
    p2.text = "Sim2Real Validation Results"
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p2.alignment = PP_ALIGN.LEFT

    # Subtitle details
    txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(3.5), Inches(8.4), Inches(3.0))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    details = [
        "Model: Llama 3.1 70B Instruct (FP8) on H100 80GB",
        "Workload: Chatbot, 10 QPS Poisson, 10K requests",
        "ISL: Gaussian(512, 200), OSL: Gaussian(256, 100)",
        "SLO: Mean TTFT <= 300 ms",
        "",
        "Tools compared: BLIS (inference-sim), AIConfigurator, llm-optimizer",
        "9 configs tested, 5 deployed successfully, 4 failed (KV cache OOM)",
    ]
    for i, line in enumerate(details):
        if i == 0:
            p = tf2.paragraphs[0]
        else:
            p = tf2.add_paragraph()
        p.text = line
        p.font.size = Pt(14)
        p.space_after = Pt(4)


def add_algorithm_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # Title
    txBox = slide.shapes.add_textbox(Inches(0.3), Inches(0.1), Inches(9.4), Inches(0.6))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "BLIS Search Algorithm: NSGA-II Multi-Objective Optimization"
    p.font.size = Pt(20)
    p.font.bold = True
    p.alignment = PP_ALIGN.LEFT

    # Left column: how it works
    txBox2 = slide.shapes.add_textbox(Inches(0.3), Inches(0.8), Inches(4.8), Inches(6.0))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True

    header = tf2.paragraphs[0]
    header.text = "How It Works"
    header.font.size = Pt(14)
    header.font.bold = True
    header.space_after = Pt(8)

    steps = [
        ("1. Generate initial population", "40 random configs from the parameter space"),
        ("2. Evaluate each config", "Run BLIS simulation (~100-250ms each, parallelized)"),
        ("3. Rank by Pareto dominance", "Non-dominated sorting into fronts"),
        ("4. Select parents", "Tournament selection favoring front rank + crowding distance"),
        ("5. Crossover + mutate", "Uniform crossover of config params, 15% per-param mutation"),
        ("6. Repeat for 4 generations", "~200 total evaluations, 3-5 seconds wall time"),
    ]
    for title, desc in steps:
        p = tf2.add_paragraph()
        p.text = title
        p.font.size = Pt(11)
        p.font.bold = True
        p.space_after = Pt(1)
        p = tf2.add_paragraph()
        p.text = f"   {desc}"
        p.font.size = Pt(10)
        p.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        p.space_after = Pt(8)

    p = tf2.add_paragraph()
    p.text = "Result: Pareto-optimal set of configs"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1F, 0x77, 0xB4)

    # Right column: parameter space + objectives
    txBox3 = slide.shapes.add_textbox(Inches(5.3), Inches(0.8), Inches(4.5), Inches(6.2))
    tf3 = txBox3.text_frame
    tf3.word_wrap = True

    header = tf3.paragraphs[0]
    header.text = "4 Objectives (simultaneous)"
    header.font.size = Pt(14)
    header.font.bold = True
    header.space_after = Pt(6)

    objectives = [
        "Maximize throughput (responses/sec)",
        "Minimize TTFT P99 (latency)",
        "Minimize GPU count (cost)",
        "Minimize KV memory (blocks)",
    ]
    for obj in objectives:
        p = tf3.add_paragraph()
        p.text = f"  {obj}"
        p.font.size = Pt(11)
        p.space_after = Pt(3)

    p = tf3.add_paragraph()
    p.text = ""
    p.space_after = Pt(10)

    p = tf3.add_paragraph()
    p.text = "6-Parameter Search Space (~1,800 configs)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.space_after = Pt(6)

    params = [
        "TP: 1, 2, 4, 8",
        "Replicas: 1..floor(8/TP)",
        "Scheduler: FCFS, SJF",
        "Max batch size: 32, 64, 128, 256, 512",
        "KV blocks: 2K, 3K, 4K, 5K, 7.5K, 10K",
        "Block size: 16, 32 tokens",
    ]
    for param in params:
        p = tf3.add_paragraph()
        p.text = f"  {param}"
        p.font.size = Pt(11)
        p.space_after = Pt(2)

    p = tf3.add_paragraph()
    p.text = ""
    p.space_after = Pt(10)

    p = tf3.add_paragraph()
    p.text = "Why NSGA-II?"
    p.font.size = Pt(14)
    p.font.bold = True
    p.space_after = Pt(6)

    reasons = [
        "Reaches 99.2% of exhaustive hypervolume at 200 evals",
        "Random search plateaus at 94.3%",
        "Crossover of Pareto parents produces Pareto offspring ~65% of the time",
    ]
    for reason in reasons:
        p = tf3.add_paragraph()
        p.text = f"  {reason}"
        p.font.size = Pt(10)
        p.space_after = Pt(3)


def add_blis_pareto_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # Title
    txBox = slide.shapes.add_textbox(Inches(0.3), Inches(0.1), Inches(9.4), Inches(0.6))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "BLIS Search Results: Pareto Frontier"
    p.font.size = Pt(20)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    # Chart image
    slide.shapes.add_picture(
        str(BLIS_PARETO_PATH), Inches(0.3), Inches(0.7), Inches(9.4), Inches(6.5)
    )


def add_chart_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # Title
    txBox = slide.shapes.add_textbox(Inches(0.3), Inches(0.1), Inches(9.4), Inches(0.6))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Pareto Front: Predicted vs Actual (TTFT vs Throughput)"
    p.font.size = Pt(20)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    # Chart image
    slide.shapes.add_picture(
        str(CHART_PATH), Inches(0.3), Inches(0.7), Inches(9.4), Inches(6.5)
    )


def add_results_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # Title
    txBox = slide.shapes.add_textbox(Inches(0.3), Inches(0.1), Inches(9.4), Inches(0.6))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Results Summary"
    p.font.size = Pt(20)
    p.font.bold = True
    p.alignment = PP_ALIGN.LEFT

    # Left column: actual results table
    txBox2 = slide.shapes.add_textbox(Inches(0.3), Inches(0.8), Inches(5.5), Inches(4.0))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True

    header = tf2.paragraphs[0]
    header.text = "Actual Results (Stable Configs)"
    header.font.size = Pt(13)
    header.font.bold = True
    header.space_after = Pt(8)

    rows = [
        ("BLIS TP4x1 (4 GPU)", "TTFT=157ms", "Max=15,056 tok/s"),
        ("BLIS TP8x1 (8 GPU)", "TTFT=197ms", "Max=13,155 tok/s"),
        ("AICfg TP4x2 (8 GPU)", "TTFT=141ms", "Max=8,188 tok/s"),
    ]
    for config, ttft, tput in rows:
        p = tf2.add_paragraph()
        p.text = f"{config}: {ttft}, {tput}"
        p.font.size = Pt(11)
        p.space_after = Pt(2)

    p = tf2.add_paragraph()
    p.text = ""
    p = tf2.add_paragraph()
    p.text = "Overloaded (batch size too small for 10 QPS):"
    p.font.size = Pt(11)
    p.font.italic = True
    p.space_after = Pt(4)

    overloaded = [
        ("AICfg TP4x1 (4 GPU, seqs=96)", "TTFT=43.6s"),
        ("llm-opt TP4x1 (4 GPU, seqs=8)", "TTFT=970s, 68% timeouts"),
    ]
    for config, note in overloaded:
        p = tf2.add_paragraph()
        p.text = f"{config}: {note}"
        p.font.size = Pt(11)
        p.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
        p.space_after = Pt(2)

    # Right column: key findings
    txBox3 = slide.shapes.add_textbox(Inches(5.8), Inches(0.8), Inches(4.0), Inches(4.5))
    tf3 = txBox3.text_frame
    tf3.word_wrap = True

    header = tf3.paragraphs[0]
    header.text = "Key Findings"
    header.font.size = Pt(13)
    header.font.bold = True
    header.space_after = Pt(8)

    findings = [
        "BLIS wins: best throughput prediction accuracy (within 15-18% of actual)",
        "Critical differentiator: batch size selection. BLIS uses max_num_seqs=512; AIConfigurator=96; llm-optimizer=8",
        "All TP<=3 configs fail (KV cache OOM with 131K context). None of the tools predicted this.",
        "BLIS underestimates TTFT by 3-5x but gets capacity right. For capacity planning, throughput accuracy matters more.",
    ]
    for finding in findings:
        p = tf3.add_paragraph()
        p.text = f"• {finding}"
        p.font.size = Pt(11)
        p.space_after = Pt(10)

    # Bottom: winner box
    txBox4 = slide.shapes.add_textbox(Inches(0.3), Inches(5.5), Inches(9.4), Inches(1.8))
    tf4 = txBox4.text_frame
    tf4.word_wrap = True
    p = tf4.paragraphs[0]
    p.text = "Winner: BLIS (inference-sim)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1F, 0x77, 0xB4)
    p.space_after = Pt(4)

    p2 = tf4.add_paragraph()
    p2.text = (
        "BLIS's binary search for max sustainable throughput produces the most operationally useful metric. "
        "Its aggressive batch sizes (512) correctly model real vLLM continuous batching behavior, "
        "while other tools' conservative settings (8-96) cause queuing failures at moderate load."
    )
    p2.font.size = Pt(11)


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)
    add_algorithm_slide(prs)
    add_blis_pareto_slide(prs)
    add_chart_slide(prs)
    add_results_slide(prs)

    out_path = OUTPUT_DIR / "config_exploration_results.pptx"
    prs.save(str(out_path))
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    main()
